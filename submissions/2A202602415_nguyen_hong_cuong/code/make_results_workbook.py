"""Build results.xlsx from the committed Colab logs and eval.py outputs."""
from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "results.xlsx"
HEADERS = {
    "Backbones": ["exp_id", "backbone", "weight_tag", "params_M", "GMAC", "resolution",
                  "epochs", "seed", "best_epoch", "macro_F1_val", "top1_val",
                  "train_s_per_epoch", "notes", "status"],
    "Training": ["exp_id", "backbone", "recipe", "seed", "macro_F1_val", "top1_val",
                 "delta_vs_T00_seed", "best_epoch", "notes", "status"],
    "Inference": ["exp_id", "method", "configuration", "seed", "views", "macro_F1_val",
                  "top1_val", "latency_batch1_ms", "notes", "status"],
    "Final": ["exp_id", "configuration", "seed", "macro_F1_val", "macro_F1_test",
              "top1_test", "balanced_acc_test", "ECE_test", "NLL_test", "status"],
    "PerClass": ["configuration", "class", "test_support", "precision_mean", "precision_std",
                 "recall_mean", "recall_std", "F1_mean", "F1_std", "status"],
    "Latency": ["configuration", "GPU", "dtype", "batch", "resolution", "mean_train_epoch_s",
                "p50_inference_ms", "p95_inference_ms", "p99_inference_ms", "notes", "status"],
    "Summary": ["rank", "exp_id", "configuration", "macro_F1_val_mean", "macro_F1_test_mean",
                "macro_F1_test_std", "top1_test_mean", "balanced_acc_test_mean", "ECE_test_mean",
                "params_M", "GMAC", "evidence_note", "status"],
}


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def format_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
    for index, column in enumerate(ws.columns, 1):
        content = [str(cell.value or "") for cell in column]
        ws.column_dimensions[get_column_letter(index)].width = min(max(max(map(len, content)) + 2, 11), 38)
    for cell in ws[1]:
        if any(key in str(cell.value) for key in ("F1", "top1", "acc")):
            letter = get_column_letter(cell.column)
            ws.conditional_formatting.add(
                f"{letter}2:{letter}{ws.max_row}",
                ColorScaleRule(start_type="min", start_color="FEE2E2", mid_type="percentile",
                               mid_value=50, mid_color="FEF3C7", end_type="max", end_color="DCFCE7"),
            )


def main() -> None:
    summaries = json.loads((ROOT / "final_summaries.json").read_text(encoding="utf-8"))
    by_key = {(row["exp_id"], int(row["seed"])): row for row in summaries}
    test_rows = {
        exp: {int(row["seed"]): row for row in read_csv(ROOT / "eval_out" / f"{exp}_per_seed.csv")}
        for exp in ("T00", "F01")
    }
    aggregate = {
        exp: json.loads((ROOT / "eval_out" / f"{exp}_summary.json").read_text(encoding="utf-8"))
        for exp in ("T00", "F01")
    }
    configs = {
        "T00": ("resnet50", "ImageNet pretrained; basic augmentation; CE; no mixing; no EMA", 64),
        "F01": ("convnext_tiny", "ImageNet pretrained; RandAugment; weighted CE; CutMix; EMA", 48),
    }

    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = {name: workbook.create_sheet(name) for name in HEADERS}
    for name, headers in HEADERS.items():
        sheets[name].append(headers)

    for exp in ("T00", "F01"):
        backbone, recipe, batch = configs[exp]
        for seed in (0, 1, 2):
            run = by_key[(exp, seed)]
            test = test_rows[exp][seed]
            sheets["Backbones"].append([
                exp, backbone, "ImageNet-1K", run["params_m"], run["gmacs_conv_linear"], 224,
                10, seed, run["best_epoch"], run["val_macro_f1"], run["val_top1"],
                run["mean_epoch_seconds"], "Full fold-0 Colab T4 run", "MEASURED",
            ])
            baseline = by_key[("T00", seed)]["val_macro_f1"]
            sheets["Training"].append([
                exp, backbone, recipe, seed, run["val_macro_f1"], run["val_top1"],
                run["val_macro_f1"] - baseline, run["best_epoch"], "Locked from validation", "MEASURED",
            ])
            sheets["Inference"].append([
                "I00", "one view FP32", exp, seed, 1, run["val_macro_f1"], run["val_top1"], "",
                "Accuracy measured; batch-1 latency was not benchmarked", "MEASURED_ACCURACY",
            ])
            sheets["Final"].append([
                exp, f"{backbone} + I00", seed, run["val_macro_f1"], float(test["macro_f1"]),
                float(test["top1"]), float(test["balanced_acc"]), float(test["ece"]), float(test["nll"]),
                "MEASURED",
            ])
        sheets["Latency"].append([
            exp, "NVIDIA T4", "AMP training / FP32 export", batch, 224,
            sum(by_key[(exp, seed)]["mean_epoch_seconds"] for seed in (0, 1, 2)) / 3,
            "", "", "", "Training time measured; inference latency not measured", "PARTIAL",
        ])

    for exp in ("T00", "F01"):
        for row in read_csv(ROOT / "eval_out" / f"{exp}_per_class.csv"):
            sheets["PerClass"].append([
                exp, row["class"], int(row["support"]), float(row["precision_mean"]),
                float(row["precision_std"]), float(row["recall_mean"]), float(row["recall_std"]),
                float(row["f1_mean"]), float(row["f1_std"]), "MEASURED_3_SEEDS",
            ])

    ranking = sorted(("T00", "F01"), key=lambda exp: aggregate[exp]["macro_f1"]["mean"], reverse=True)
    for rank, exp in enumerate(ranking, 1):
        backbone, recipe, _ = configs[exp]
        agg = aggregate[exp]
        val_mean = sum(by_key[(exp, seed)]["val_macro_f1"] for seed in (0, 1, 2)) / 3
        run0 = by_key[(exp, 0)]
        sheets["Summary"].append([
            rank, exp, f"{backbone}: {recipe}", val_mean, agg["macro_f1"]["mean"],
            agg["macro_f1"]["std"], agg["top1"]["mean"], agg["balanced_acc"]["mean"],
            agg["ece"]["mean"], run0["params_m"], run0["gmacs_conv_linear"],
            "eval.py over three independent seeds", "MEASURED",
        ])

    for sheet in sheets.values():
        format_sheet(sheet)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    checked = load_workbook(OUTPUT, read_only=True, data_only=False)
    assert checked.sheetnames == list(HEADERS)
    checked.close()
    print(OUTPUT)


if __name__ == "__main__":
    main()
